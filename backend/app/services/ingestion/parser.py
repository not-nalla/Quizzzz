import os
import logging
from typing import List, Dict, Any

logger = logging.getLogger("syllabusiq.ingestion.parser")

class DocumentParser:
    @staticmethod
    def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
        pages = []
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text").strip()
                if text:
                    pages.append({"page": page_num + 1, "text": text})
            doc.close()
        except Exception as e:
            logger.error(f"Error parsing PDF {file_path}: {e}")
            raise
        return pages

    @staticmethod
    def parse_docx(file_path: str) -> List[Dict[str, Any]]:
        pages = []
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            
            # Approximate pages by paragraph blocks (approx 4 paragraphs per page)
            joined_paragraphs = "\n\n".join(full_text)
            chunks_approx = [p for p in joined_paragraphs.split("\n\n") if p]
            page_size = 4
            for i in range(0, max(len(chunks_approx), 1), page_size):
                page_text = "\n\n".join(chunks_approx[i:i + page_size])
                pages.append({"page": (i // page_size) + 1, "text": page_text})
        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path}: {e}")
            raise
        return pages

    @staticmethod
    def parse_pptx(file_path: str) -> List[Dict[str, Any]]:
        pages = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide_num, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    pages.append({"page": slide_num + 1, "text": "\n".join(slide_text)})
        except Exception as e:
            logger.error(f"Error parsing PPTX {file_path}: {e}")
            raise
        return pages

    @staticmethod
    def parse_txt(file_path: str) -> List[Dict[str, Any]]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read().strip()
            return [{"page": 1, "text": content}]
        except Exception as e:
            logger.error(f"Error parsing TXT {file_path}: {e}")
            raise

    @classmethod
    def parse(cls, file_path: str, filename: str) -> List[Dict[str, Any]]:
        lower = filename.lower()
        if lower.endswith(".pdf"):
            return cls.parse_pdf(file_path)
        elif lower.endswith(".docx"):
            return cls.parse_docx(file_path)
        elif lower.endswith(".pptx"):
            return cls.parse_pptx(file_path)
        elif lower.endswith(".txt"):
            return cls.parse_txt(file_path)
        else:
            # Default to text reader
            return cls.parse_txt(file_path)
